#!/usr/bin/env python3

import json
import re
import sys
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import xlrd


KNOWN_HEADERS = {
    "customer_name": {"单位名称", "单位", "客户名称", "公司名称"},
    "project_name": {"工程项目", "项目名称"},
    "project_belonging": {"项目归属"},
    "receivable_amount": {"对账金额", "已对账金额"},
    "contract_receivable": {"按合同应收金额", "合同应收金额", "按合同应收"},
    "received_amount": {"实际收款金额", "已收金额", "收款金额"},
    "unpaid_amount": {"按合同未收款项（元）", "未收金额", "未收款项", "未收款"},
    "cumulative_receivable": {"累计应收（元）", "累计应收", "应收账款余额"},
    "statement_date": {"对账时间", "对账日期", "对账日"},
    "received_date": {"回款日期", "实际收款日期", "收款日期"},
    "business_type": {"业务类型", "业务定义", "款项类别", "业务类别"},
    "time_node": {"时间节点", "时点"},
    "remark": {"备注"},
    "current_status": {"目前状况", "当前状况"},
    "credit_status": {"资信状况"},
}

LEDGER_HEADERS = {
    "statement_date": {"对账时间", "对账日期", "对账日"},
    "quantity": {"方量", "数量"},
    "statement_amount": {"对账金额", "已对账金额"},
    "contract_receivable": {"按合同应收金额", "合同应收金额"},
    "received_date": {"实际收款时间", "实际收款日期", "收款时间", "收款日期"},
    "received_amount": {"实际收款金额", "已收金额", "收款金额"},
    "cumulative_receivable": {"累计应收", "累计应收（元）", "应收账款余额"},
    "invoice_date": {"开票时间", "开票日期"},
    "invoice_quantity": {"开票方量"},
    "invoice_amount": {"开票金额"},
    "remark": {"备注"},
}

SKIP_KEYWORDS = ("合计", "小计", "总计", "汇总")
HEADER_LIKE_CUSTOMERS = {"单位", "单位名称", "客户名称", "公司名称", "购货单位", "结算单位"}
HEADER_LIKE_PROJECTS = {"工程项目", "项目名称", "项目", "工程名称"}
NOISE_CUSTOMER_NAMES = set(SKIP_KEYWORDS) | HEADER_LIKE_CUSTOMERS | {"待补数据"}
NOISE_PROJECT_NAMES = set(SKIP_KEYWORDS) | HEADER_LIKE_PROJECTS | {"待补数据", "合计/待补数据"}

def normalize_text(value):
    if value is None:
        return None
    text = str(value).replace("\u3000", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text or None


def normalize_customer_name(value):
    text = normalize_text(value)
    if text is None:
        return None
    # Excel 总表中存在“客户名 + 汇总”的口径，这类文本应映射回真实客户名，
    # 以免把汇总展示名称误当成新客户。
    text = re.sub(r"\s*汇总\s*$", "", text)
    return text or None


def is_noise_customer_name(value):
    text = normalize_customer_name(value)
    if text is None:
        return True
    return text in NOISE_CUSTOMER_NAMES


def is_noise_project_name(value):
    text = normalize_text(value)
    if text is None:
        return True
    if text in NOISE_PROJECT_NAMES:
        return True
    return bool(re.search(r"汇总\s*$", text))


def normalize_header(value):
    text = normalize_text(value)
    if text is None:
        return None
    return re.sub(r"[\s_（）()\-]+", "", text).lower()


NORMALIZED_HEADER_MAP = {
    key: {normalize_header(item) for item in values if normalize_header(item)}
    for key, values in KNOWN_HEADERS.items()
}

NORMALIZED_LEDGER_HEADER_MAP = {
    key: {normalize_header(item) for item in values if normalize_header(item)}
    for key, values in LEDGER_HEADERS.items()
}


def to_decimal(value):
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = normalize_text(value)
    if text is None:
        return None
    text = text.replace(",", "").replace("，", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def to_number(value):
    decimal_value = to_decimal(value)
    return float(decimal_value) if decimal_value is not None else None


def to_json_safe(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return serialize_date_or_text(value)
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (int, float, str, bool)):
        return value
    return normalize_text(value)


def serialize_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = normalize_text(value)
    if text is None:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m", "%Y/%m", "%Y年%m月%d日", "%Y年%m月"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.date().isoformat()
        except ValueError:
            continue
    return None


def serialize_date_or_text(value):
    return serialize_date(value) or normalize_text(value)


def detect_header_index(rows):
    for index, row in enumerate(rows):
        normalized = {normalize_header(cell) for cell in row if normalize_header(cell)}
        matched = sum(
            1
            for values in NORMALIZED_HEADER_MAP.values()
            if any(
                cell == value or value in cell or cell in value
                for cell in normalized
                for value in values
            )
        )
        if matched >= 2:
            return index
    return None


def rows_from_xlsx(path):
    workbook = load_workbook(path, data_only=True)
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        raw_rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
        yield sheet_name, raw_rows, ws


def rows_from_xls(path):
    book = xlrd.open_workbook(path, formatting_info=False)
    for sheet_name in book.sheet_names():
        sheet = book.sheet_by_name(sheet_name)
        raw_rows = []
        for row_index in range(sheet.nrows):
            values = []
            for col in range(sheet.ncols):
                cell = sheet.cell(row_index, col)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = datetime(*xlrd.xldate_as_tuple(value, book.datemode)).date().isoformat()
                values.append(value)
            raw_rows.append(values)
        yield sheet_name, raw_rows, sheet


def normalize_business_type(value):
    text = normalize_text(value)
    if not text:
        return ""
    if text in {"诉", "已诉", "诉讼"}:
        return "已诉"
    if text == "待起诉":
        return "待起诉"
    if text == "尾款":
        return "尾款"
    return text


def to_english_key(header):
    normalized = normalize_header(header) or "extra_field"
    normalized = re.sub(r"[^a-z0-9]+", "_", normalized).strip("_")
    if not normalized:
        normalized = "extra_field"
    if re.match(r"^\d", normalized):
        normalized = f"field_{normalized}"
    return normalized


def is_summary_like(*values):
    for value in values:
        text = normalize_text(value) or ""
        if any(keyword in text for keyword in SKIP_KEYWORDS):
            return True
    return False


def is_customer_profile_sheet(raw_rows):
    if len(raw_rows) < 4:
        return False
    first = normalize_text(raw_rows[0][0] if raw_rows[0] else None) or ""
    second = normalize_text(raw_rows[1][0] if len(raw_rows) > 1 and raw_rows[1] else None) or ""
    third = normalize_text(raw_rows[2][0] if len(raw_rows) > 2 and raw_rows[2] else None) or ""
    return first.startswith("客户名称：") and second.startswith("工程项目：") and third.startswith("合同结算条款：")


def extract_label_value(raw_value, label):
    text = normalize_text(raw_value) or ""
    prefix = f"{label}："
    if text.startswith(prefix):
        return text[len(prefix):].strip() or None
    return None


def is_ledger_header_row(row):
    normalized = {normalize_header(cell) for cell in row if normalize_header(cell)}
    matched = 0
    for values in NORMALIZED_LEDGER_HEADER_MAP.values():
        if any(
            cell == value or value in cell or cell in value
            for cell in normalized
            for value in values
        ):
            matched += 1
    return matched >= 3


def map_ledger_row(headers, row):
    source = {}
    matched_headers = set()
    for header, value in zip(headers, row):
        key = normalize_text(header)
        if key:
            source[key] = value

    def pick(field_name):
        normalized_targets = NORMALIZED_LEDGER_HEADER_MAP[field_name]
        for header_name, value in source.items():
            normalized_header_name = normalize_header(header_name)
            if normalized_header_name and any(
                normalized_header_name == target
                or target in normalized_header_name
                or normalized_header_name in target
                for target in normalized_targets
            ):
                matched_headers.add(header_name)
                return value
        return None

    mapped = {
        "statementDate": serialize_date_or_text(pick("statement_date")),
        "quantity": to_number(pick("quantity")),
        "statementAmount": to_number(pick("statement_amount")),
        "contractReceivable": to_number(pick("contract_receivable")),
        "receivedDate": serialize_date_or_text(pick("received_date")),
        "receivedAmount": to_number(pick("received_amount")),
        "cumulativeReceivable": to_number(pick("cumulative_receivable")),
        "invoiceDate": serialize_date_or_text(pick("invoice_date")),
        "invoiceQuantity": to_number(pick("invoice_quantity")),
        "invoiceAmount": to_number(pick("invoice_amount")),
        "remark": normalize_text(pick("remark")) or "",
    }

    extra_fields = {}
    for header_name, value in source.items():
        if header_name in matched_headers:
            continue
        normalized_value = normalize_text(value)
        if normalized_value in (None, "") and not isinstance(value, (int, float)):
            continue
        extra_fields[to_english_key(header_name)] = to_json_safe(value)

    if extra_fields:
        mapped["extraFields"] = extra_fields

    if not any(
        mapped.get(key) not in (None, "", 0)
        for key in [
            "statementDate",
            "quantity",
            "statementAmount",
            "contractReceivable",
            "receivedDate",
            "receivedAmount",
            "cumulativeReceivable",
            "invoiceDate",
            "invoiceQuantity",
            "invoiceAmount",
            "remark",
        ]
    ):
        return None

    return mapped


def extract_summary_amount(row, label):
    for index, value in enumerate(row):
        text = normalize_text(value) or ""
        if label in text:
            for next_value in row[index + 1 :]:
                amount = to_number(next_value)
                if amount is not None:
                    return amount
    return None


def parse_profile_sheet(sheet_name, raw_rows):
    customer_name = normalize_customer_name(
        extract_label_value(raw_rows[0][0] if raw_rows[0] else None, "客户名称")
    )
    project_name = extract_label_value(raw_rows[1][0] if len(raw_rows) > 1 and raw_rows[1] else None, "工程项目")
    settlement_clause = extract_label_value(raw_rows[2][0] if len(raw_rows) > 2 and raw_rows[2] else None, "合同结算条款")
    headers = raw_rows[3] if len(raw_rows) > 3 else []

    rows = []
    summary = {
        "contractUnpaidTotal": None,
        "uninvoicedTotal": None,
        "extraSummary": {},
    }

    for row_index, row in enumerate(raw_rows[4:], start=5):
        if not any(normalize_text(cell) for cell in row):
            continue
        if is_ledger_header_row(row):
            continue

        contract_unpaid = extract_summary_amount(row, "按合同未收款")
        uninvoiced = extract_summary_amount(row, "未制票")
        if contract_unpaid is not None:
            summary["contractUnpaidTotal"] = contract_unpaid
        if uninvoiced is not None:
            summary["uninvoicedTotal"] = uninvoiced
        if contract_unpaid is not None or uninvoiced is not None:
            for col_index, value in enumerate(row, start=1):
                text = normalize_text(value)
                if text and any(keyword in text for keyword in ("按合同未收款", "未制票", "合计", "验证")):
                    summary["extraSummary"][f"col_{col_index}"] = text
            continue

        if any((normalize_text(cell) or "") in SKIP_KEYWORDS for cell in row):
            continue

        mapped = map_ledger_row(headers, row)
        if not mapped:
            continue
        rows.append(
            {
                "rowOrder": len(rows) + 1,
                "sourceRowNumber": row_index,
                **mapped,
            }
        )

    if not customer_name:
        return None

    if not customer_name and not project_name and not rows:
        return None

    return {
        "sheetName": sheet_name,
        "customerName": customer_name,
        "projectName": project_name,
        "settlementClause": settlement_clause,
        "summary": {
            "contractUnpaidTotal": summary["contractUnpaidTotal"],
            "uninvoicedTotal": summary["uninvoicedTotal"],
            "extraSummary": summary["extraSummary"] or {},
        },
        "rows": rows,
    }


def cell_to_sheet_text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return serialize_date_or_text(value)
    if isinstance(value, Decimal):
        return str(float(value))
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value)


def width_to_pixels(width):
    if width is None:
        return None
    return max(60, int(float(width) * 7 + 12))


def build_sheet_view_data(sheet_name, raw_rows, sheet_obj=None):
    last_row = 0
    last_col = 0
    for r_index, row in enumerate(raw_rows, start=1):
        has_value = False
        for c_index, value in enumerate(row, start=1):
            if value not in (None, ""):
                has_value = True
                last_col = max(last_col, c_index)
        if has_value:
            last_row = r_index

    if last_row == 0 or last_col == 0:
        return {
            "name": sheet_name,
            "freeze": "A5",
            "rows": {},
            "cols": {},
            "merges": [],
        }

    rows = {}
    for ri in range(last_row):
        cells = {}
        current_row = raw_rows[ri] if ri < len(raw_rows) else []
        for ci in range(last_col):
            value = current_row[ci] if ci < len(current_row) else None
            text = cell_to_sheet_text(value)
            if text is None or text == "":
                continue
            cells[ci] = {"text": text}

        row_meta = {}
        if sheet_obj is not None:
            height = getattr(sheet_obj.row_dimensions.get(ri + 1), "height", None)
            if height:
                row_meta["height"] = int(height)
        if cells:
            row_meta["cells"] = cells
        if row_meta:
            rows[ri] = row_meta

    cols = {}
    if sheet_obj is not None:
        for ci in range(last_col):
            letter = get_column_letter(ci + 1)
            width = getattr(sheet_obj.column_dimensions.get(letter), "width", None)
            px = width_to_pixels(width)
            if px:
                cols[ci] = {"width": px}

    merges = []
    if sheet_obj is not None and getattr(sheet_obj, "merged_cells", None):
        merges = [str(rng) for rng in sheet_obj.merged_cells.ranges]
    elif hasattr(sheet_obj, "merged_cells"):
        for rlo, rhi, clo, chi in getattr(sheet_obj, "merged_cells", []):
            start = f"{get_column_letter(clo + 1)}{rlo + 1}"
            end = f"{get_column_letter(chi)}{rhi}"
            merges.append(f"{start}:{end}")

    return {
        "name": sheet_name,
        "freeze": "A5",
        "rows": rows,
        "cols": cols,
        "merges": merges,
    }


def map_row(headers, row):
    source = {}
    matched_headers = set()
    for header, value in zip(headers, row):
        key = normalize_text(header)
        if key:
            source[key] = value

    def pick(field_name):
        normalized_targets = NORMALIZED_HEADER_MAP[field_name]
        for header_name, value in source.items():
            normalized_header_name = normalize_header(header_name)
            if normalized_header_name and any(
                normalized_header_name == target
                or target in normalized_header_name
                or normalized_header_name in target
                for target in normalized_targets
            ):
                matched_headers.add(header_name)
                return value
        return None

    customer_name = normalize_customer_name(pick("customer_name"))
    project_name = normalize_text(pick("project_name"))
    if not customer_name and not project_name:
        return None, source

    remark_parts = []
    base_remark = normalize_text(pick("remark"))
    current_status = normalize_text(pick("current_status"))
    if base_remark:
        remark_parts.append(base_remark)
    if current_status and current_status not in remark_parts:
        remark_parts.append(current_status)

    mapped = {
        "customerName": customer_name,
        "creditStatus": normalize_text(pick("credit_status")),
        "currentStatus": current_status,
        "projectName": project_name,
        "projectBelonging": normalize_text(pick("project_belonging")),
        "projectStatus": current_status or normalize_business_type(pick("business_type")) or "进行中",
        "receivableAmount": to_number(pick("receivable_amount")) or 0,
        "contractReceivable": to_number(pick("contract_receivable")),
        "receivedAmount": to_number(pick("received_amount")) or 0,
        "unpaidAmount": to_number(pick("unpaid_amount")),
        "cumulativeReceivable": to_number(pick("cumulative_receivable")),
        "statementDate": serialize_date(pick("statement_date")),
        "receivedDate": serialize_date(pick("received_date")),
        "businessType": normalize_business_type(pick("business_type")),
        "timeNode": normalize_text(pick("time_node")),
        "remark": "；".join(remark_parts),
    }

    if mapped["contractReceivable"] is None:
        mapped["contractReceivable"] = mapped["receivableAmount"]
    if mapped["unpaidAmount"] is None:
        mapped["unpaidAmount"] = max(mapped["contractReceivable"] - mapped["receivedAmount"], 0)
    if mapped["cumulativeReceivable"] is None:
        mapped["cumulativeReceivable"] = mapped["contractReceivable"]
    if not mapped["timeNode"]:
        mapped["timeNode"] = mapped["statementDate"][:7] if mapped["statementDate"] else ""

    extra_fields = {}
    for header_name, value in source.items():
        if header_name in matched_headers:
            continue
        normalized_value = normalize_text(value)
        if normalized_value in (None, "") and not isinstance(value, (int, float)):
            continue
        extra_fields[to_english_key(header_name)] = to_json_safe(value)

    if extra_fields:
        mapped["extraFields"] = extra_fields

    return mapped, source


def build_risk(row):
    business_type = row.get("businessType", "")
    current_status = row.get("currentStatus", "") or ""
    remark = row.get("remark", "") or ""
    text = f"{remark} {current_status}".strip()
    if business_type in {"待起诉", "已诉"}:
        return {
            "customerName": row["customerName"],
            "projectName": row.get("projectName"),
            "riskType": "诉讼",
            "riskLevel": "高",
            "source": "Excel导入",
            "remark": remark or business_type,
        }
    if any(keyword in current_status for keyword in ["停工", "资金紧张", "未回款"]):
        if "停工" in current_status:
            risk_type = "经营异常"
        elif "资金紧张" in current_status:
            risk_type = "财务异常"
        else:
            risk_type = "付款逾期"
        return {
            "customerName": row["customerName"],
            "projectName": row.get("projectName"),
            "riskType": risk_type,
            "riskLevel": "中",
            "source": "Excel导入",
            "remark": remark or current_status,
        }
    if any(keyword in remark for keyword in ["风险", "无法支付"]):
        risk_type = "财务异常" if "无法支付" in remark else "沟通异常"
        return {
            "customerName": row["customerName"],
            "projectName": row.get("projectName"),
            "riskType": risk_type,
            "riskLevel": "中",
            "source": "Excel导入",
            "remark": remark,
        }
    return None


def parse_workbook(path):
    path = Path(path)
    loader = rows_from_xlsx if path.suffix.lower() == ".xlsx" else rows_from_xls
    customers = OrderedDict()
    projects = OrderedDict()
    receivables = []
    risks = OrderedDict()
    mapped_fields = set()
    unmapped_columns = set()
    errors = []
    skipped_rows = 0
    ledger_profiles = []

    for sheet_name, raw_rows, sheet_obj in loader(path):
        if is_customer_profile_sheet(raw_rows):
            profile = parse_profile_sheet(sheet_name, raw_rows)
            if profile:
                profile["sheetData"] = build_sheet_view_data(sheet_name, raw_rows, sheet_obj)
                profile["workbookSource"] = path.name
                ledger_profiles.append(profile)

        header_index = detect_header_index(raw_rows)
        if header_index is None:
            continue
        headers = raw_rows[header_index]
        for header in headers:
            header_text = normalize_text(header)
            if not header_text:
                continue
            normalized_header_text = normalize_header(header_text)
            if any(
                normalized_header_text == item
                or item in normalized_header_text
                or normalized_header_text in item
                for items in NORMALIZED_HEADER_MAP.values()
                for item in items
            ):
                mapped_fields.add(header_text)
            else:
                unmapped_columns.add(header_text)

        for row_offset, row in enumerate(raw_rows[header_index + 1 :], start=header_index + 2):
            if not any(normalize_text(cell) for cell in row):
                continue
            mapped, _source = map_row(headers, row)
            if not mapped:
                continue
            if is_noise_customer_name(mapped.get("customerName")) or is_noise_project_name(mapped.get("projectName")):
                skipped_rows += 1
                continue
            if is_summary_like(mapped.get("customerName"), mapped.get("projectName"), _source.get("单位"), _source.get("公司名称")):
                errors_is_zero_amount = to_number(_source.get("对账金额")) or to_number(_source.get("已对账金额")) or 0
                if not mapped.get("customerName") or not mapped.get("projectName") or errors_is_zero_amount == 0:
                    skipped_rows += 1
                    continue
            if not mapped["customerName"]:
                amount = (
                    mapped.get("receivableAmount")
                    or mapped.get("contractReceivable")
                    or mapped.get("cumulativeReceivable")
                    or 0
                )
                if amount == 0:
                    skipped_rows += 1
                    continue
                errors.append({
                    "rowRef": f"{sheet_name}:{row_offset}",
                    "message": "客户名称为空",
                    "rawPayload": {normalize_text(headers[i]) or f"col_{i}": to_json_safe(row[i]) for i in range(min(len(headers), len(row)))},
                })
                continue
            if not mapped.get("projectName"):
                if is_summary_like(_source.get("工程项目"), _source.get("项目名称"), _source.get("单位"), _source.get("公司名称")):
                    skipped_rows += 1
                    continue
                errors.append({
                    "rowRef": f"{sheet_name}:{row_offset}",
                    "message": "项目名称为空",
                    "rawPayload": {normalize_text(headers[i]) or f"col_{i}": to_json_safe(row[i]) for i in range(min(len(headers), len(row)))},
                })
                continue

            existing_customer = customers.get(mapped["customerName"])
            if not existing_customer:
                customers[mapped["customerName"]] = {
                    "customerName": mapped["customerName"],
                    "creditStatus": mapped.get("creditStatus"),
                    "currentStatus": mapped.get("currentStatus"),
                    "remark": mapped.get("remark", ""),
                }
            else:
                if not existing_customer.get("creditStatus") and mapped.get("creditStatus"):
                    existing_customer["creditStatus"] = mapped.get("creditStatus")
                if not existing_customer.get("currentStatus") and mapped.get("currentStatus"):
                    existing_customer["currentStatus"] = mapped.get("currentStatus")
                if not existing_customer.get("remark") and mapped.get("remark"):
                    existing_customer["remark"] = mapped.get("remark", "")

            if mapped.get("projectName") and not is_noise_project_name(mapped.get("projectName")):
                project_key = f"{mapped['customerName']}::{mapped['projectName']}"
                existing_project = projects.get(project_key)
                if not existing_project:
                    projects[project_key] = {
                        "customerName": mapped["customerName"],
                        "projectName": mapped["projectName"],
                        "projectBelonging": mapped.get("projectBelonging"),
                        "projectStatus": mapped.get("projectStatus") or "进行中",
                        "creditStatus": mapped.get("creditStatus"),
                        "currentStatus": mapped.get("currentStatus"),
                        "remark": mapped.get("remark", ""),
                    }
                else:
                    if not existing_project.get("projectBelonging") and mapped.get("projectBelonging"):
                        existing_project["projectBelonging"] = mapped.get("projectBelonging")
                    if not existing_project.get("creditStatus") and mapped.get("creditStatus"):
                        existing_project["creditStatus"] = mapped.get("creditStatus")
                    if not existing_project.get("currentStatus") and mapped.get("currentStatus"):
                        existing_project["currentStatus"] = mapped.get("currentStatus")
                    if (
                        existing_project.get("projectStatus") in (None, "", "进行中")
                        and mapped.get("projectStatus")
                        and mapped.get("projectStatus") != "进行中"
                    ):
                        existing_project["projectStatus"] = mapped.get("projectStatus")
                    if not existing_project.get("remark") and mapped.get("remark"):
                        existing_project["remark"] = mapped.get("remark", "")

            receivables.append(
                {
                    "customerName": mapped["customerName"],
                    "projectName": mapped.get("projectName"),
                    "receivableAmount": mapped.get("receivableAmount") or 0,
                    "contractReceivable": mapped.get("contractReceivable") or 0,
                    "receivedAmount": mapped.get("receivedAmount") or 0,
                    "unpaidAmount": mapped.get("unpaidAmount") or 0,
                    "cumulativeReceivable": mapped.get("cumulativeReceivable") or 0,
                    "statementDate": mapped.get("statementDate"),
                    "receivedDate": mapped.get("receivedDate"),
                    "businessType": mapped.get("businessType") or "",
                    "timeNode": mapped.get("timeNode") or sheet_name,
                    "projectBelonging": mapped.get("projectBelonging"),
                    "creditStatus": mapped.get("creditStatus"),
                    "currentStatus": mapped.get("currentStatus"),
                    "remark": mapped.get("remark", ""),
                    "extraFields": mapped.get("extraFields", {}),
                }
            )

            risk = build_risk(mapped)
            if risk:
                key = f"{risk['customerName']}::{risk.get('projectName') or ''}::{risk['riskType']}"
                risks.setdefault(key, risk)

    return {
        "customers": list(customers.values()),
        "projects": list(projects.values()),
        "receivables": receivables,
        "risks": list(risks.values()),
        "mappedFields": sorted(mapped_fields),
        "unmappedColumns": sorted(unmapped_columns),
        "errors": errors,
        "skippedRows": skipped_rows,
        "ledgerProfiles": ledger_profiles,
    }


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "missing file path"}, ensure_ascii=False))
        sys.exit(1)

    payload = parse_workbook(sys.argv[1])
    print(json.dumps(payload, ensure_ascii=False))


if __name__ == "__main__":
    main()
