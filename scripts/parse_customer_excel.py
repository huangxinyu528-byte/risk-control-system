#!/usr/bin/env python3

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
import xlrd


KNOWN_HEADERS = {
    "customer_name": {"客户名称", "公司名称", "单位名称", "单位", "主体名称"},
    "customer_code": {"客户编号", "编号", "客户代码", "客户编码"},
    "sales_owner_name": {"业务负责人", "负责人", "销售负责人", "业务员", "责任人"},
    "unified_social_credit_code": {"统一社会信用代码", "社会信用代码", "信用代码"},
    "industry": {"行业", "所属行业"},
    "contact_name": {"联系人", "联系人姓名"},
    "contact_phone": {"联系电话", "手机号", "电话", "联系方式"},
    "credit_status": {"资信状况", "信用状况"},
    "current_status": {"目前状况", "当前状况"},
    "customer_nature": {"客户性质", "主体性质", "客户类型"},
    "remark": {"备注", "说明"},
    "credit_limit": {"信用额度", "授信额度"},
}

SKIP_KEYWORDS = {"合计", "小计", "总计", "汇总", "待补数据", "单位", "单位名称", "客户名称", "公司名称", "购货单位", "结算单位"}


def normalize_text(value):
    if value is None:
        return None
    text = str(value).replace("\u3000", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text or None


def normalize_header(value):
    text = normalize_text(value)
    if text is None:
        return None
    return re.sub(r"[\s_（）()\-]+", "", text).lower()


def normalize_customer_name(value):
    text = normalize_text(value)
    if text is None:
        return None
    text = re.sub(r"\s*汇总\s*$", "", text)
    return text or None


def is_noise_customer_name(value):
    text = normalize_customer_name(value)
    if text is None:
        return True
    return text in SKIP_KEYWORDS


NORMALIZED_HEADER_MAP = {
    key: {normalize_header(item) for item in values if normalize_header(item)}
    for key, values in KNOWN_HEADERS.items()
}


def to_json_safe(value):
    if value is None:
        return None
    if isinstance(value, (int, float, str, bool)):
        return value
    return normalize_text(value)


def detect_header_index(rows):
    for index, row in enumerate(rows):
        normalized = {normalize_header(cell) for cell in row if normalize_header(cell)}
        if not normalized:
            continue
        matched = sum(
            1
            for values in NORMALIZED_HEADER_MAP.values()
            if any(
                cell == value or value in cell or cell in value
                for cell in normalized
                for value in values
            )
        )
        if matched >= 2 and any(
            any(
                cell == value or value in cell or cell in value
                for cell in normalized
                for value in NORMALIZED_HEADER_MAP["customer_name"]
            )
            for _ in [0]
        ):
            return index
    return None


def rows_from_xlsx(path):
    workbook = load_workbook(path, data_only=True)
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        raw_rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
        yield sheet_name, raw_rows


def rows_from_xls(path):
    book = xlrd.open_workbook(path, formatting_info=False)
    for sheet_name in book.sheet_names():
        sheet = book.sheet_by_name(sheet_name)
        raw_rows = []
        for row_index in range(sheet.nrows):
            raw_rows.append(sheet.row_values(row_index))
        yield sheet_name, raw_rows


def map_header_to_key(header):
    normalized = normalize_header(header)
    if not normalized:
        return None

    for key, values in NORMALIZED_HEADER_MAP.items():
        if normalized in values:
            return key

    for key, values in NORMALIZED_HEADER_MAP.items():
        if any(value in normalized for value in values if len(value) >= 4):
            return key
    return None


def parse_file(path):
    loader = rows_from_xlsx if path.suffix.lower() == ".xlsx" else rows_from_xls
    customers = []
    mapped_fields = set()
    unmapped_columns = set()
    errors = []
    skipped_rows = 0

    for sheet_name, rows in loader(path):
        header_index = detect_header_index(rows)
        if header_index is None:
            continue

        headers = rows[header_index]
        mapped_headers = {}
        for col_index, header in enumerate(headers):
            key = map_header_to_key(header)
            if key:
                mapped_headers[col_index] = key
                mapped_fields.add(key)
            else:
                text = normalize_text(header)
                if text:
                    unmapped_columns.add(text)

        for row_offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            raw_payload = {}
            payload = {}
            for col_index, cell in enumerate(row):
                header_text = normalize_text(headers[col_index] if col_index < len(headers) else None)
                if header_text:
                    raw_payload[header_text] = to_json_safe(cell)
                if col_index in mapped_headers:
                    payload[mapped_headers[col_index]] = normalize_text(cell)

            if not any(normalize_text(cell) for cell in row):
                continue

            customer_name = normalize_customer_name(payload.get("customer_name"))
            if not customer_name:
                errors.append(
                    {
                        "rowRef": f"{sheet_name}!{row_offset}",
                        "message": "缺少客户名称",
                        "rawPayload": raw_payload,
                    }
                )
                continue

            if is_noise_customer_name(customer_name):
                skipped_rows += 1
                continue

            customers.append(
                {
                    "customerName": customer_name,
                    "customerCode": payload.get("customer_code") or None,
                    "salesOwnerName": payload.get("sales_owner_name") or None,
                    "unifiedSocialCreditCode": payload.get("unified_social_credit_code") or None,
                    "industry": payload.get("industry") or None,
                    "contactName": payload.get("contact_name") or None,
                    "contactPhone": payload.get("contact_phone") or None,
                    "creditStatus": payload.get("credit_status") or None,
                    "currentStatus": payload.get("current_status") or None,
                    "customerNature": payload.get("customer_nature") or None,
                    "remark": payload.get("remark") or None,
                    "creditLimit": payload.get("credit_limit") or None,
                }
            )

    return {
        "customers": customers,
        "mappedFields": sorted(mapped_fields),
        "unmappedColumns": sorted(unmapped_columns),
        "errors": errors,
        "skippedRows": skipped_rows,
    }


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: parse_customer_excel.py <excel-path>")

    path = Path(sys.argv[1]).expanduser().resolve()
    if not path.exists():
        raise SystemExit(f"Excel file not found: {path}")

    result = parse_file(path)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
