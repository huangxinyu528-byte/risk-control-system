#!/usr/bin/env python3

import argparse
import hashlib
import json
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
import xlrd


DEFAULT_TEMPLATE_PATH = Path("data/众程建材客户风险管理系统模板_最终版(2).xlsx")
DEFAULT_RECEIVABLE_PATH = Path("data/众程建材应收账款分类表2026.2.28(2)(1).xls")


def normalize_text(value):
    if value is None:
        return None
    text = str(value).replace("\u3000", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text or None


def to_decimal(value):
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = normalize_text(value)
    if text is None:
        return None
    text = text.replace(",", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def decimal_str(value):
    if value is None:
        return None
    quantized = value.quantize(Decimal("0.01"))
    return format(quantized, "f")


def to_int(value):
    decimal_value = to_decimal(value)
    if decimal_value is None:
        return None
    return int(decimal_value)


def to_bool(value):
    text = normalize_text(value)
    if text is None:
        return None
    if text in {"是", "Y", "y", "yes", "true", "True", "1"}:
        return True
    if text in {"否", "N", "n", "no", "false", "False", "0"}:
        return False
    return None


def serialize_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = normalize_text(value)
    if text is None:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m", "%Y/%m"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.date().isoformat()
        except ValueError:
            continue
    return None


def stable_code(prefix, *parts):
    raw = "|".join([normalize_text(part) or "" for part in parts])
    digest = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12].upper()
    return f"{prefix}{digest}"


def load_sheet_rows(ws):
    headers = None
    rows = []
    for row in ws.iter_rows(values_only=True):
        values = [normalize_text(cell) for cell in row]
        if not any(values):
            continue
        if headers is None:
            headers = values
            continue
        item = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            item[header] = row[index] if index < len(row) else None
        if any(value not in (None, "") for value in item.values()):
            rows.append(item)
    return rows


def load_template_data(template_path):
    workbook = load_workbook(template_path, data_only=True)
    params = {}
    if "Params" in workbook.sheetnames:
        for row in load_sheet_rows(workbook["Params"]):
            key = normalize_text(row.get("参数项"))
            if key:
                params[key] = row.get("参数值")

    payload = {
        "params": params,
        "customers": [],
        "contracts": [],
        "risks": [],
        "scores": [],
    }

    if "Customers" in workbook.sheetnames:
        payload["customers"] = load_sheet_rows(workbook["Customers"])
    if "Contracts" in workbook.sheetnames:
        payload["contracts"] = load_sheet_rows(workbook["Contracts"])
    if "Risks" in workbook.sheetnames:
        payload["risks"] = load_sheet_rows(workbook["Risks"])
    if "Scores" in workbook.sheetnames:
        payload["scores"] = load_sheet_rows(workbook["Scores"])

    return payload


def parse_xls_sheet(sheet):
    header_row = None
    headers = []
    rows = []
    for row_index in range(sheet.nrows):
        row_values = [normalize_text(sheet.cell_value(row_index, col)) for col in range(sheet.ncols)]
        if header_row is None and "序号" in row_values:
            header_row = row_index
            headers = row_values
            continue
        if header_row is None:
            continue
        if not any(row_values):
            continue
        item = {}
        for col_index, header in enumerate(headers):
            if not header:
                continue
            cell = sheet.cell(row_index, col_index)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = datetime(*xlrd.xldate_as_tuple(value, sheet.book.datemode)).date().isoformat()
            item[header] = value
        rows.append(item)
    return rows


def parse_receivable_workbook(receivable_path):
    book = xlrd.open_workbook(receivable_path, formatting_info=False)

    main_sheet_name = next((name for name in book.sheet_names() if "汇总" in name), book.sheet_names()[0])
    litigation_sheet_name = next((name for name in book.sheet_names() if name == "Sheet1"), None)
    history_sheet_name = next((name for name in book.sheet_names() if name == "2.15"), None)

    main_rows = parse_xls_sheet(book.sheet_by_name(main_sheet_name))
    litigation_rows = parse_xls_sheet(book.sheet_by_name(litigation_sheet_name)) if litigation_sheet_name else []
    history_rows = parse_xls_sheet(book.sheet_by_name(history_sheet_name)) if history_sheet_name else []

    return {
        "main_sheet_name": main_sheet_name,
        "main_rows": main_rows,
        "litigation_sheet_name": litigation_sheet_name,
        "litigation_rows": litigation_rows,
        "history_sheet_name": history_sheet_name,
        "history_rows": history_rows,
    }


def map_contract_status(value):
    text = normalize_text(value)
    mapping = {
        "执行中": "active",
        "进行中": "active",
        "待回款": "active",
        "已完成": "completed",
        "已终止": "terminated",
        "终止": "terminated",
        "暂停中": "suspended",
        "暂停": "suspended",
        "草稿": "draft",
    }
    return mapping.get(text, "active")


def map_risk_source(value):
    text = normalize_text(value)
    mapping = {
        "人工录入": "manual",
        "手工录入": "manual",
        "系统计算": "system",
        "系统导入": "system",
        "天眼查同步": "tianyancha",
    }
    return mapping.get(text, "system")


def map_risk_level(value):
    text = normalize_text(value)
    mapping = {
        "低": "low",
        "中": "medium",
        "高": "high",
        "重大": "critical",
        "严重": "critical",
    }
    return mapping.get(text, "high")


def score_level(score):
    if score >= 90:
        return "A+"
    if score >= 80:
        return "A"
    if score >= 70:
        return "B"
    if score >= 60:
        return "C"
    return "D"


def detect_import_month(main_sheet_name):
    match = re.search(r"(\d{1,2})\.(\d{1,2})", main_sheet_name or "")
    if match:
        month = int(match.group(1))
        return f"2026-{month:02d}"
    return date.today().strftime("%Y-%m")


def build_payload(template_data, receivable_data):
    params = template_data["params"]
    import_month = detect_import_month(receivable_data["main_sheet_name"])
    customers = {}
    contracts = {}
    receivables = []
    risks = {}
    history_map = {}

    for row in receivable_data["history_rows"]:
        company = normalize_text(row.get("单位"))
        project_name = normalize_text(row.get("工程项目"))
        if not company or not project_name:
            continue
        history_map[(company, project_name)] = {
            "monthlyReceivedAmount": decimal_str(to_decimal(row.get("2月收款额"))),
            "balanceAtMonthStart": decimal_str(to_decimal(row.get("应收账款余额"))),
        }

    template_customer_rows = template_data["customers"]
    for row in template_customer_rows:
        name = normalize_text(row.get("客户名称"))
        if not name:
            continue
        customer_id = normalize_text(row.get("客户ID")) or stable_code("ZC-KH-", name)
        customers[name] = {
            "customerId": customer_id,
            "name": name,
            "unifiedSocialCreditCode": normalize_text(row.get("统一社会信用代码")),
            "industry": normalize_text(row.get("行业")),
            "contactName": normalize_text(row.get("联系人")),
            "contactPhone": normalize_text(row.get("联系电话")),
            "salesOwnerName": normalize_text(row.get("业务负责人")),
            "currentCreditLevel": normalize_text(row.get("当前信用等级")),
            "creditLimit": decimal_str(to_decimal(row.get("信用额度(元)"))),
            "stopSupply": to_bool(row.get("是否停供")),
            "remark": normalize_text(row.get("备注")),
            "source": "template",
        }

    current_company = None
    scanned_main_rows = 0
    skipped_subtotal_rows = 0

    for row in receivable_data["main_rows"]:
        company = normalize_text(row.get("单位")) or current_company
        project_name = normalize_text(row.get("工程项目"))
        if company:
            current_company = company
        if not company or not project_name:
            continue
        scanned_main_rows += 1
        if any(keyword in project_name for keyword in ("小计", "合计")):
            skipped_subtotal_rows += 1
            continue

        customer = customers.get(company)
        if customer is None:
            customer = {
                "customerId": stable_code("ZC-KH-", company),
                "name": company,
                "unifiedSocialCreditCode": None,
                "industry": None,
                "contactName": None,
                "contactPhone": None,
                "salesOwnerName": None,
                "currentCreditLevel": None,
                "creditLimit": None,
                "stopSupply": False,
                "remark": "由应收账款明细自动生成，待补充客户主数据。",
                "source": "receivable-sheet",
            }
            customers[company] = customer

        contract_no = stable_code("AUTO-CT-", company, project_name)
        contracts[contract_no] = {
            "contractNo": contract_no,
            "customerRef": customer["customerId"],
            "title": project_name,
            "contractBusinessId": None,
            "contractStatus": "active",
            "currency": "CNY",
            "amount": decimal_str(to_decimal(row.get("按合同应收金额"))),
            "signedAt": None,
            "billingPeriodDays": None,
            "salesOwnerName": customer.get("salesOwnerName"),
            "startDate": None,
            "endDate": None,
            "paymentTerms": None,
            "invoiceRequirement": None,
            "remark": "由应收账款明细自动生成，账期和签订信息待补充。",
            "summary": f"自动导入合同占位：{company} / {project_name}",
            "openclawSummary": f"{company} 项目 {project_name} 的占位合同记录。",
            "source": "receivable-sheet",
        }

        monthly_values = history_map.get((company, project_name), {})
        amount = to_decimal(row.get("按合同应收金额")) or to_decimal(row.get("累计应收（元）")) or Decimal("0")
        received_amount = to_decimal(row.get("实际收款金额")) or Decimal("0")
        unpaid_amount = to_decimal(row.get("按合同未收款项（元）"))
        if unpaid_amount is None:
            unpaid_amount = amount - received_amount

        receivable_no = stable_code(
            "AUTO-AR-",
            receivable_data["main_sheet_name"],
            company,
            project_name,
            row.get("按合同应收金额"),
            row.get("实际收款金额"),
        )
        receivables.append(
            {
                "receivableNo": receivable_no,
                "customerRef": customer["customerId"],
                "contractRef": contract_no,
                "projectName": project_name,
                "amount": decimal_str(amount),
                "receivedAmount": decimal_str(received_amount),
                "unpaidAmount": decimal_str(unpaid_amount),
                "monthlyReceivedAmount": monthly_values.get("monthlyReceivedAmount"),
                "salesOwnerName": customer.get("salesOwnerName"),
                "billingPeriodDays": None,
                "dueDate": None,
                "overdueDays": None,
                "agingBucket": None,
                "collectionLevel": None,
                "riskStatus": normalize_text(row.get("款项类别")) or normalize_text(row.get("业务定义")),
                "receivableStatus": "completed" if unpaid_amount <= 0 else "pending",
                "lastCollectionAt": None,
                "expectedPaymentDate": None,
                "openclawSummary": f"{company} - {project_name} 当前未收 {decimal_str(unpaid_amount) or '0.00'} 元。",
                "receivableSnapshot": {
                    "sourceSheet": receivable_data["main_sheet_name"],
                    "businessDefinition": normalize_text(row.get("业务定义")),
                    "paymentCategory": normalize_text(row.get("款项类别")),
                    "notReconciledAmount": decimal_str(to_decimal(row.get("未对账金额"))),
                    "totalReceivable": decimal_str(to_decimal(row.get("累计应收（元）"))),
                    "paymentToOutputRatio": str(row.get("收款占总产值比")) if row.get("收款占总产值比") not in (None, "") else None,
                    "paymentToProgressRatio": str(row.get("收款占进度款比例")) if row.get("收款占进度款比例") not in (None, "") else None,
                    "originalRow": {k: (normalize_text(v) if not isinstance(v, (int, float)) else v) for k, v in row.items()},
                },
            }
        )

        payment_category = normalize_text(row.get("款项类别"))
        if payment_category == "诉讼":
            risk_code = stable_code("AUTO-RISK-", company, project_name, "诉讼", receivable_data["main_sheet_name"])
            risks[risk_code] = {
                "riskCode": risk_code,
                "customerRef": customer["customerId"],
                "riskType": "诉讼",
                "riskLevel": "high",
                "source": "system",
                "title": f"{company} 存在诉讼相关应收项目",
                "description": f"项目 {project_name} 在应收账款分类表中标记为诉讼款项。",
                "handled": False,
                "eventAt": None,
                "detectedAt": None,
                "resolvedAt": None,
                "status": "open",
                "sourceDetail": receivable_data["main_sheet_name"],
                "remark": "由应收账款分类表自动生成。",
                "openclawSummary": f"{company} 的项目 {project_name} 存在诉讼风险。",
                "payload": {
                    "projectName": project_name,
                    "unpaidAmount": decimal_str(unpaid_amount),
                    "sourceSheet": receivable_data["main_sheet_name"],
                },
            }

    for row in template_data["contracts"]:
        customer_name = normalize_text(row.get("客户名称"))
        project_name = normalize_text(row.get("项目名称"))
        if not customer_name or not project_name:
            continue
        customer = customers.get(customer_name)
        if customer is None:
            customer = {
                "customerId": normalize_text(row.get("客户ID")) or stable_code("ZC-KH-", customer_name),
                "name": customer_name,
                "unifiedSocialCreditCode": None,
                "industry": None,
                "contactName": None,
                "contactPhone": None,
                "salesOwnerName": normalize_text(row.get("业务负责人")),
                "currentCreditLevel": None,
                "creditLimit": None,
                "stopSupply": False,
                "remark": "由合同模板补充生成。",
                "source": "template-contract",
            }
            customers[customer_name] = customer
        contract_no = normalize_text(row.get("合同编号")) or stable_code("AUTO-CT-", customer_name, project_name)
        contracts[contract_no] = {
            "contractNo": contract_no,
            "customerRef": customer["customerId"],
            "title": project_name,
            "contractBusinessId": normalize_text(row.get("合同ID")),
            "contractStatus": map_contract_status(row.get("项目状态")),
            "currency": "CNY",
            "amount": decimal_str(to_decimal(row.get("合同金额(元)"))),
            "signedAt": serialize_date(row.get("签订日期")),
            "billingPeriodDays": to_int(row.get("账期(天)")),
            "salesOwnerName": normalize_text(row.get("业务负责人")),
            "startDate": None,
            "endDate": serialize_date(row.get("预计到期日")),
            "paymentTerms": None,
            "invoiceRequirement": None,
            "remark": None,
            "summary": f"{customer_name} / {project_name}",
            "openclawSummary": f"{customer_name} 的项目 {project_name} 合同信息。",
            "source": "template-contract",
        }

    for row in template_data["risks"]:
        customer_name = normalize_text(row.get("客户名称"))
        risk_type = normalize_text(row.get("风险类型"))
        if not customer_name or not risk_type:
            continue
        customer = customers.get(customer_name)
        if customer is None:
            customer = {
                "customerId": normalize_text(row.get("客户ID")) or stable_code("ZC-KH-", customer_name),
                "name": customer_name,
                "stopSupply": False,
                "remark": "由风险模板补充生成。",
                "source": "template-risk",
            }
            customers[customer_name] = customer
        risk_code = normalize_text(row.get("风险ID")) or stable_code("AUTO-RISK-", customer_name, risk_type, row.get("风险日期"))
        risks[risk_code] = {
            "riskCode": risk_code,
            "customerRef": customer["customerId"],
            "riskType": risk_type,
            "riskLevel": map_risk_level(row.get("风险等级")),
            "source": map_risk_source(row.get("风险来源")),
            "title": f"{customer_name} - {risk_type}",
            "description": normalize_text(row.get("处理动作")),
            "handled": to_bool(row.get("是否已处理")) or False,
            "eventAt": serialize_date(row.get("风险日期")),
            "detectedAt": None,
            "resolvedAt": None,
            "status": "closed" if (to_bool(row.get("是否已处理")) is True) else "open",
            "sourceDetail": normalize_text(row.get("风险来源")),
            "remark": normalize_text(row.get("备注")),
            "openclawSummary": f"{customer_name} 存在 {risk_type} 风险。",
            "payload": {"source": "template-risk"},
        }

    for row in receivable_data["litigation_rows"]:
        customer_name = normalize_text(row.get("公司名称"))
        project_name = normalize_text(row.get("项目名称"))
        amount = to_decimal(row.get("金额"))
        if not customer_name or not project_name or "小计" in project_name:
            continue
        customer = customers.get(customer_name)
        if customer is None:
            customer = {
                "customerId": stable_code("ZC-KH-", customer_name),
                "name": customer_name,
                "stopSupply": False,
                "remark": "由诉讼客户清单自动生成。",
                "source": "litigation-sheet",
            }
            customers[customer_name] = customer
        risk_code = stable_code("AUTO-RISK-", customer_name, project_name, "诉讼清单")
        risks[risk_code] = {
            "riskCode": risk_code,
            "customerRef": customer["customerId"],
            "riskType": "诉讼",
            "riskLevel": "high",
            "source": "system",
            "title": f"{customer_name} 诉讼客户清单命中",
            "description": f"项目 {project_name} 出现在诉讼客户应收款一览表中。",
            "handled": False,
            "eventAt": None,
            "detectedAt": None,
            "resolvedAt": None,
            "status": "open",
            "sourceDetail": receivable_data["litigation_sheet_name"],
            "remark": "由诉讼客户应收款一览表自动生成。",
            "openclawSummary": f"{customer_name} 的项目 {project_name} 在诉讼客户清单中，金额 {decimal_str(amount) or '0.00'} 元。",
            "payload": {
                "projectName": project_name,
                "amount": decimal_str(amount),
                "sourceSheet": receivable_data["litigation_sheet_name"],
            },
        }

    scores = {}
    template_scores = template_data["scores"]
    for row in template_scores:
        customer_name = normalize_text(row.get("客户名称"))
        if not customer_name:
            continue
        customer = customers.get(customer_name)
        if customer is None:
            customer = {
                "customerId": normalize_text(row.get("客户ID")) or stable_code("ZC-KH-", customer_name),
                "name": customer_name,
                "stopSupply": False,
                "remark": "由评分模板补充生成。",
                "source": "template-score",
            }
            customers[customer_name] = customer
        score_month = serialize_date(row.get("评分日期"))
        if score_month:
            score_month = score_month[:7]
        else:
            score_month = import_month
        total_score = to_int(row.get("总分"))
        credit_risk_score = to_int(row.get("信用风险得分(40)")) or 0
        business_risk_score = to_int(row.get("经营风险得分(30)")) or 0
        financial_risk_score = to_int(row.get("财务风险得分(30)")) or 0
        if total_score is None:
            total_score = credit_risk_score + business_risk_score + financial_risk_score
        scores[(customer["customerId"], score_month)] = {
            "customerRef": customer["customerId"],
            "scoreMonth": score_month,
            "creditRiskScore": credit_risk_score,
            "businessRiskScore": business_risk_score,
            "financialRiskScore": financial_risk_score,
            "score": total_score,
            "scoreLevel": normalize_text(row.get("信用等级")) or score_level(total_score),
            "scoreVersion": "import-v1",
            "scoreReason": normalize_text(row.get("评语")) or "由评分模板导入。",
            "keyAttention": (normalize_text(row.get("信用等级")) in {"C", "D"}) if normalize_text(row.get("信用等级")) else total_score < 70,
            "riskFlags": {
                "lawsuitCount": to_int(row.get("诉讼件数")) or 0,
                "executedCount": to_int(row.get("被执行次数")) or 0,
                "dishonesty": to_bool(row.get("失信标记")) or False,
                "taxAbnormal": to_bool(row.get("税务异常标记")) or False,
                "historyOverdueCount": to_int(row.get("历史逾期次数")) or 0,
                "maxOverdueDays": to_int(row.get("最大逾期天数")),
            },
            "openclawSummary": f"{customer_name} 在 {score_month} 的信用等级为 {normalize_text(row.get('信用等级')) or score_level(total_score)}。",
            "source": "template-score",
        }

    customer_metrics = defaultdict(lambda: {"amount": Decimal("0"), "received": Decimal("0"), "lawsuitCount": 0})
    for receivable in receivables:
        metric = customer_metrics[receivable["customerRef"]]
        metric["amount"] += to_decimal(receivable["amount"]) or Decimal("0")
        metric["received"] += to_decimal(receivable["receivedAmount"]) or Decimal("0")
    for risk in risks.values():
        if risk["riskType"] == "诉讼":
            customer_metrics[risk["customerRef"]]["lawsuitCount"] += 1

    lawsuit_penalty = to_int(params.get("诉讼单件扣分")) or 10
    executed_penalty = to_int(params.get("被执行单件扣分")) or 15
    dishonesty_penalty = to_int(params.get("失信标记扣分")) or 20
    tax_penalty = to_int(params.get("税务异常扣分")) or 10

    for customer in customers.values():
        key = (customer["customerId"], import_month)
        if key in scores:
            continue
        metric = customer_metrics[customer["customerId"]]
        total_amount = metric["amount"]
        received_amount = metric["received"]
        payment_rate = Decimal("0")
        if total_amount > 0:
            payment_rate = received_amount / total_amount
        credit_risk_score = max(0, 40 - (metric["lawsuitCount"] * lawsuit_penalty))
        business_risk_score = max(0, min(30, int((payment_rate * Decimal("30")).quantize(Decimal("1")))))
        financial_risk_score = 30
        credit_limit = to_decimal(customer.get("creditLimit"))
        if credit_limit and total_amount > 0:
            utilization = total_amount / credit_limit
            if utilization > Decimal("1"):
                financial_risk_score = 10
            elif utilization > Decimal("0.9"):
                financial_risk_score = 18
            elif utilization >= Decimal("0.7"):
                financial_risk_score = 24
        total_score = credit_risk_score + business_risk_score + financial_risk_score
        scores[key] = {
            "customerRef": customer["customerId"],
            "scoreMonth": import_month,
            "creditRiskScore": credit_risk_score,
            "businessRiskScore": business_risk_score,
            "financialRiskScore": financial_risk_score,
            "score": total_score,
            "scoreLevel": score_level(total_score),
            "scoreVersion": "import-v1",
            "scoreReason": "根据导入应收与风险数据自动生成首版评分；账期、到期日等缺失字段未参与扣分。",
            "keyAttention": total_score < 70 or metric["lawsuitCount"] > 0,
            "riskFlags": {
                "lawsuitCount": metric["lawsuitCount"],
                "executedCount": 0,
                "dishonesty": False,
                "taxAbnormal": False,
                "historyOverdueCount": 0,
                "maxOverdueDays": None,
                "assumptions": [
                    "账期和到期日缺失时不参与自动预警。",
                    "经营风险使用收款比例进行简化换算。",
                    "财务风险优先使用信用额度占用率；缺失额度时按默认值处理。",
                ],
            },
            "openclawSummary": f"{customer['name']} 在 {import_month} 的自动评分为 {total_score} 分，等级 {score_level(total_score)}。",
            "source": "computed",
        }

    return {
        "meta": {
            "importMonth": import_month,
            "sourceFiles": {
                "template": str(DEFAULT_TEMPLATE_PATH),
                "receivable": str(DEFAULT_RECEIVABLE_PATH),
            },
            "assumptions": [
                "第一版主应收数据使用 2.28汇总 工作表。",
                "2.15 工作表仅用于补充本月回款字段。",
                "账期、到期日为空时保留为空，不触发错误，也不参与自动预警。",
                "缺失客户主数据时，按公司名称生成稳定 customerId。",
                "缺失合同编号时，按客户名称 + 项目名称生成稳定 contractNo。",
                "缺失评分明细时，基于应收和风险记录生成首版自动评分。",
            ],
        },
        "params": params,
        "customers": list(customers.values()),
        "contracts": list(contracts.values()),
        "accountsReceivable": receivables,
        "riskRecords": list(risks.values()),
        "creditScores": list(scores.values()),
        "summary": {
            "mainRowsScanned": scanned_main_rows,
            "subtotalRowsSkipped": skipped_subtotal_rows,
            "templateCustomers": len(template_customer_rows),
            "templateContracts": len(template_data["contracts"]),
            "templateRisks": len(template_data["risks"]),
            "templateScores": len(template_scores),
            "litigationRows": len(receivable_data["litigation_rows"]),
            "customersPrepared": len(customers),
            "contractsPrepared": len(contracts),
            "receivablesPrepared": len(receivables),
            "risksPrepared": len(risks),
            "scoresPrepared": len(scores),
        },
    }


def main():
    parser = argparse.ArgumentParser(description="提取众程建材 Excel 数据并映射为结构化 JSON。")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE_PATH))
    parser.add_argument("--receivable", default=str(DEFAULT_RECEIVABLE_PATH))
    args = parser.parse_args()

    template_path = Path(args.template)
    receivable_path = Path(args.receivable)

    template_data = load_template_data(template_path)
    receivable_data = parse_receivable_workbook(receivable_path)
    payload = build_payload(template_data, receivable_data)
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
